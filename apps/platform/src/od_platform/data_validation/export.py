"""CSV / Excel 返工清单导出。

提供标注团队可直接打开的整改清单:
    - CSV: 标准库 csv 模块，零依赖
    - Excel: openpyxl 软依赖，未安装时自动降级为 CSV

每行 = 一个可整改的问题条目。
列: check_name, severity, split, file_path, line_number, issue_type, description, suggested_action
"""

from __future__ import annotations

import csv
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

_REWORK_COLUMNS = [
    "check_name",
    "severity",
    "split",
    "file_path",
    "line_number",
    "issue_type",
    "description",
    "suggested_action",
]


def write_rework_csv(report: object, output_path: Path) -> Path:
    """导出 CSV 返工清单。使用标准库 csv，零额外依赖。"""
    rows = _collect_rework_rows(report)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with open(output_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=_REWORK_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)

    logger.info("CSV 返工清单已写入: %s (%d 行)", output_path, len(rows))
    return output_path


def write_rework_excel(report: object, output_path: Path) -> Path:
    """导出 Excel 返工清单。软依赖 openpyxl，未安装时自动降级为 CSV。

    每个 check 一个 sheet (sheet 名 = check_name)，便于标注团队按问题类型分批处理。
    """
    try:
        import openpyxl
    except ImportError:
        logger.info("openpyxl 未安装，降级为 CSV 格式")
        return write_rework_csv(report, output_path.with_suffix(".csv"))

    rows = _collect_rework_rows(report)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # 删除默认空 sheet

    # 按 check_name 分组
    by_check: dict[str, list[dict]] = {}
    for row in rows:
        by_check.setdefault(row["check_name"], []).append(row)

    for check_name, check_rows in sorted(by_check.items()):
        ws = wb.create_sheet(title=check_name[:31])  # sheet 名最长 31 字符
        ws.append(_REWORK_COLUMNS)
        for row in check_rows:
            ws.append([row.get(col, "") for col in _REWORK_COLUMNS])
        # 设置列宽
        for col_idx in range(1, len(_REWORK_COLUMNS) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 18

    wb.save(output_path)
    logger.info("Excel 返工清单已写入: %s (%d 行, %d sheets)", output_path, len(rows), len(by_check))
    return output_path


def _collect_rework_rows(report: object) -> list[dict[str, str]]:
    """从 report.results 中提取所有可整改问题条目。"""
    rows: list[dict[str, str]] = []

    for r in report.results:
        if r.severity not in ("WARNING", "ERROR"):
            continue

        det = r.details
        action = det.get("action", "人工复核")

        if r.name == "yaml_schema" and "problems" in det:
            for p in det["problems"]:
                rows.append({
                    "check_name": r.name,
                    "severity": r.severity,
                    "split": "",
                    "file_path": str(report.yaml_path),
                    "line_number": "",
                    "issue_type": "yaml_config",
                    "description": p,
                    "suggested_action": action,
                })

        elif r.name == "pair_existence":
            missing_examples = det.get("missing_examples", {})
            for split, paths in missing_examples.items():
                for path in paths:
                    rows.append({
                        "check_name": r.name,
                        "severity": r.severity,
                        "split": split,
                        "file_path": path,
                        "line_number": "",
                        "issue_type": "missing_label",
                        "description": "缺少同名 label 文件",
                        "suggested_action": action,
                    })

        elif r.name == "label_format":
            for e in det.get("errors_preview", []):
                rows.append({
                    "check_name": r.name,
                    "severity": r.severity,
                    "split": "",
                    "file_path": e.get("label", ""),
                    "line_number": str(e.get("line_no", "")),
                    "issue_type": e.get("kind", "format_error"),
                    "description": e.get("detail", ""),
                    "suggested_action": action,
                })

        elif r.name == "split_uniqueness":
            for o in det.get("overlaps", []):
                for stem in o.get("preview", []):
                    rows.append({
                        "check_name": r.name,
                        "severity": r.severity,
                        "split": f'{o["split_a"]}->{o["split_b"]}',
                        "file_path": stem,
                        "line_number": "",
                        "issue_type": "split_leakage",
                        "description": f"图像 stem 同时出现在 {o['split_a']} 和 {o['split_b']}",
                        "suggested_action": action,
                    })
        else:
            if "reason" in det or "problems" in det:
                rows.append({
                    "check_name": r.name,
                    "severity": r.severity,
                    "split": "",
                    "file_path": "",
                    "line_number": "",
                    "issue_type": det.get("reason", "unknown"),
                    "description": r.summary,
                    "suggested_action": action,
                })

    return rows
